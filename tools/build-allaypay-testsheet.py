# -*- coding: utf-8 -*-
"""Build the AllayPay/Finagy ACH certification test sheet.

Only the Comment column is filled in, per Andrii Seniv's instruction
("You need to fill the Comment column only - add transaction IDs and simple
explanation"). Answers and Production testing are left for Finagy.

Every result below came from a real run against https://token-staging.finagy.com
between 2026-09-09 and 2026-09-11 (UTC), merchant user 251. The transaction ids
are real transactions on that environment and can be looked up on your side.

2026-09-11: revised after Finagy settled 639245589443869934 and
639245635306201861 and returned 639245942322405169 (R01). SALE-05 and SALE-06
were re-run against the settled pair; SALE-08..10 now carry real rows.

2026-09-11 (evening, UTC): revised after Andrii Seniv answered the questions -
the refunded debit now reads status 40 and the credit status 4, pageId is a
cursor that ends at null, tranStatus is defined, the reserve is sandbox data.
Every row was re-read live at 15:52 UTC before rewriting it.

Run: python tools/build-allaypay-testsheet.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "ALLAYPAY-ACH-TEST-CASES.xlsx"

HEADERS = [
    "Test Case ID",
    "Test Case",
    "Test Procedure",
    "User Experience Recommendation",
    "Comment",
    "Answers",
    "Production testing",
]

# id, test case, procedure, UX recommendation, comment
ROWS = [
    (
        "SALE-01",
        "Initiate a valid ACH API Transaction",
        "Initiate a payment for any amount less than $1",
        "Transaction result is displayed to a client",
        "PASS.\n"
        "Transaction ID (authorizationId): 639245942322405169\n"
        "Our reference (uniqueTranId): ENL-CERT-S01B  |  Amount: $0.50  |  Status: 2 (pending)\n"
        "\n"
        "What we did: paid $0.50 through the ACH Hosted Payment Page, signing in to the chime "
        "sandbox account with Bank Connect (Ribbit).\n"
        "What came back on the redirect: code 200, payment_id 639245942322405169, "
        "status PENDING, merchant_token MT-48abc5d9-8f8f-4fd8-8e6d-fe88b5b6357a.\n"
        "Confirmed server-side with POST /api/echeck/retrieve: status 2 (pending), $0.50, "
        "account ...8563.\n"
        "\n"
        "The buyer sees the result on our checkout page. We do NOT treat the redirect as "
        "payment - the order stays 'pending' until the poller sees status 16 (settled). "
        "Nothing ships before that.\n"
        "\n"
        "Later: this transaction was returned R01 on 2026-09-10 (status 8), as you reported. "
        "On a return our poller moves the order to 'returned', puts the stock and loyalty "
        "points back, and emails the buyer and our staff.",
    ),
    (
        "SALE-02",
        "Initiate an invalid ACH API Transaction",
        "Initiate a payment for any amount more than $1000 which results in an error response",
        "Transaction result is displayed to a client",
        "PASS.\n"
        "Our reference (uniqueTranId): ENL-CERT-S02  |  Amount: $1,500.00  |  No transaction created\n"
        "\n"
        "What we did: repeated SALE-01 with the amount set to $1,500.00, same bank account.\n"
        "What came back: your page showed the modal 'ACH Transaction Result - Transaction "
        "Declined. Please use another account or click Cancel to return to payment selection.'\n"
        "POST /api/echeck/retrieve on ENL-CERT-S02 returns not found, so no transaction was "
        "created and nothing was debited. Correct behaviour.\n"
        "\n"
        "The error is shown to the buyer and the order is left unpaid.\n"
        "One request: the decline message has no reason code, so we cannot tell an "
        "over-limit decline from a bad-account decline. If there is a code we can read, we "
        "would like to show buyers something more useful than 'Transaction Declined.'",
    ),
    (
        "SALE-03",
        "Update ACH API Transaction Status",
        "Receive an up-to-date transaction status",
        "No customer facing UX is expected",
        "PASS.\n"
        "Endpoint: POST /api/echeck/retrieve\n"
        "\n"
        "Works with either key, both verified live:\n"
        "  - by your authorizationId 639245942322405169 -> status 2 (pending)\n"
        "  - by our uniqueTranId ENL-CERT-S01B -> the same transaction\n"
        "\n"
        "A real status change was read back end to end: 639245946244424782 was seen at "
        "status 2 (pending), voided, then re-read at status 32 (voided).\n"
        "\n"
        "This is our only source of truth, because there is no webhook. "
        "POST /api/ach/poll runs on a 30-minute cron, re-reads every open transaction and "
        "updates the order. Our mapping: 16 = paid; 8 or 24 = returned; 1 or 32 = cancelled; "
        "40 (refunded) = cancelled; 2 and 4 = still waiting. Status 24 (late return) is flagged "
        "to a human and does not auto-restock, because the goods may already have shipped.\n"
        "Staff can also ask on demand for any single order (\"Check with Finagy\" in our admin "
        "console) - that is how a refund or return on an order already marked paid reaches it.",
    ),
    (
        "SALE-04",
        "Void an ACH Transaction",
        "Initiate an ACH Transaction. Void it before transaction is settled",
        "No customer facing UX is expected",
        "PASS.\n"
        "Transaction ID: 639245946244424782  |  Our reference: ENL-CERT-S04B  |  Amount: $0.50\n"
        "Endpoint: POST /api/echeck/void\n"
        "\n"
        "What we did: created the transaction the same way as SALE-01 (redirect returned "
        "payment_id 639245946244424782, status PENDING), confirmed it was at status 2, then "
        "voided it immediately.\n"
        "What came back: {\"successful\":true}. Re-reading the transaction shows "
        "status 32 (voided).\n"
        "\n"
        "Staff reversals are a Void / Refund button on the order in our admin console, backed by "
        "one endpoint that reads the live status first, voids at status 2, refunds at status 16 "
        "and refuses anything else. A bank order can no longer be cancelled in our books while "
        "your status says the debit is still live - staff are sent to Void instead.",
    ),
    (
        "SALE-05",
        "Fail to void an ACH Transaction",
        "Initiate an ACH Transaction. Void it after transaction is settled",
        "No customer facing UX is expected",
        "PASS (on a settled transaction).\n"
        "Transaction ID: 639245635306201861  |  Our reference: ENL-MTU8KE4Z  |  Amount: $0.22  |  "
        "Settled 2026-09-10\n"
        "Endpoint: POST /api/echeck/void\n"
        "\n"
        "What we did: confirmed status 16 (settled), then called void on 2026-09-11 02:58 UTC.\n"
        "What came back: HTTP 400, {\"successful\":false,\"message\":\"Cannot void. Transaction is "
        "already processed\"}\n"
        "Re-read afterwards: still status 16. The failed void changed nothing.\n"
        "\n"
        "Also refused, earlier, on an already-voided transaction 639245946244424782: HTTP 400, "
        "\"Transaction was already voided.\"\n"
        "\n"
        "We treat this as a normal outcome, not a crash: the message is shown to the staff "
        "member, the order is left as it was, and no stock or loyalty points are released.",
    ),
    (
        "SALE-06",
        "Refund an ACH Transaction",
        "Initiate an ACH Transaction. Refund it after transaction is settled",
        "No customer facing UX is expected",
        "PASS.\n"
        "Transaction ID: 639245589443869934  |  Our reference: ENL-MTU5U6TG  |  Amount: $0.22  |  "
        "Settled 2026-09-10\n"
        "Refund credit created: 639246923014123850\n"
        "Endpoint: POST /api/echeck/refund\n"
        "\n"
        "What we did: confirmed status 16 (settled), then refunded it on 2026-09-11 02:58 UTC "
        "(after your message saying no refund was visible yet).\n"
        "What came back: {\"successful\":true,\"message\":\"\",\"authorizationId\":"
        "\"639246923014123850\"}\n"
        "Re-read after your settlement correction (2026-09-11 15:52 UTC):\n"
        "  - original debit 639245589443869934: status 40 (refunded)\n"
        "  - refund credit 639246923014123850: tranCode 2 (credit), $0.22, same account "
        "...4906, status 4 (sent to bank). As you explained, it will show in querysettlements "
        "once it settles.\n"
        "\n"
        "Why not 639245942322405169, the one we had set aside for this: it was returned R01, "
        "so there was nothing to refund. We used a settled one, as you suggested.\n"
        "\n"
        "Understood already: your refund raises a credit for the FULL original amount, so a "
        "partial refund has to be a manual credit in your portal.\n"
        "\n"
        "How we handle it: our refund button reads the live status first and only refunds at "
        "status 16, so a debit at 40 can never be refunded twice. The credit's authorizationId "
        "is stored on the order, so the credit is matched to it when it settles.\n"
        "\n"
        "QUESTION (still open): the credit has no uniqueTranId. Can a refund carry our "
        "reference, or is there a field that links the credit back to the original debit?",
    ),
    (
        "SALE-07",
        "Fail to refund an ACH Transaction",
        "Initiate an ACH Transaction. Refund it before transaction is settled",
        "No customer facing UX is expected",
        "PASS.\n"
        "Transaction ID: 639245942322405169  |  Our reference: ENL-CERT-S01B  |  Status at the "
        "time: 2 (pending)\n"
        "Endpoint: POST /api/echeck/refund\n"
        "\n"
        "What we did: attempted a refund on a transaction that had not settled yet - exactly "
        "the procedure in this row.\n"
        "What came back: HTTP 400, {\"successful\":false,\"message\":\"Failed to refund the "
        "transaction. Transaction waiting to be processed or was sent to bank.\"}\n"
        "Re-read afterwards: still status 2. The failed refund changed nothing, which is what "
        "we want.\n"
        "\n"
        "A second refusal, on a voided transaction 639245946244424782: HTTP 400, "
        "\"Transaction was already voided.\"\n"
        "\n"
        "Handled as a normal outcome: the message is shown to staff, the order is not marked "
        "refunded, and nothing is restocked.",
    ),
    (
        "SALE-08",
        "QuerySettlements",
        "Pull up-to-date Settlements information",
        "No customer facing UX is expected",
        "PASS.\n"
        "Endpoint: POST /api/echeck/querysettlements\n"
        "Request: start 2026-09-01, end 2026-09-13 (re-run 2026-09-11 15:52 UTC)\n"
        "Response: HTTP 200, 3 rows:\n"
        "  - 639245589443869934  $0.22  transactionCode D  settleDate 2026-09-10\n"
        "  - 639245635306201861  $0.22  transactionCode D  settleDate 2026-09-10\n"
        "  - 639245942322405169  $0.50  settleDate null  returnDate 2026-09-10  R01\n"
        "The refund credit 639246923014123850 is not listed yet because it is still in flight "
        "(status 4), as you explained.\n"
        "\n"
        "How we use it: shown in our admin console for any date range, every page read, each "
        "row matched to its order (refund credits included) - with excludeReturnedItems=true, "
        "since returns come from queryreturns. We match rows on authorizationId, because in "
        "this feed uniqueTranId holds your authorizationId rather than our reference.\n"
        "\n"
        "Paging, understood: pageId is a Base64 cursor naming the last row, and the report ends "
        "when the next request returns pageId null. With excludeReturnedItems=true it works "
        "exactly that way: page 1 = 2 rows, page 2 = 0 rows and pageId null.\n"
        "\n"
        "ONE ISSUE TO REPORT: with excludeReturnedItems=false the report never reaches null.\n"
        "  - page 1: 3 rows, pageId = {\"SettlementId\":25285,\"ReturnId\":0,\"ReserveId\":0}\n"
        "  - page 2 (sending that cursor): 1 row - the returned 639245942322405169 again - and "
        "pageId = {\"SettlementId\":0,\"ReturnId\":0,\"ReserveId\":0}\n"
        "    (raw: eyJTZXR0bGVtZW50SWQiOjAsIlJldHVybklkIjowLCJSZXNlcnZlSWQiOjB9)\n"
        "  - page 3 (sending that cursor): the same 3 rows as page 1, and page 1's cursor again\n"
        "The cursor does not seem to record the returned row's ReturnId, so a client that "
        "follows it until null loops forever. We guard against it on our side (we stop when a "
        "cursor repeats), but you may want to look at it.",
    ),
    (
        "SALE-09",
        "QueryReturns",
        "Pull up-to-date Returns information",
        "No customer facing UX is expected",
        "PASS.\n"
        "Endpoint: POST /api/echeck/queryreturns\n"
        "Request: start 2026-09-01, end 2026-09-13 (re-run 2026-09-11 15:52 UTC)\n"
        "Response: HTTP 200, 1 row:\n"
        "  - 639245942322405169  $0.50  dateReturned 2026-09-10  returnReason R01  "
        "returnAmount 0.50  tranStatus 1\n"
        "Page 2 (sending cursor {\"ReturnId\":21457}): 0 rows, pageId null - end of report.\n"
        "\n"
        "Thank you for defining tranStatus (1 = settled, debited back next business day; "
        "2 = returned while pending deposit; 3 = Notification of Change, information only). "
        "How we use this report:\n"
        "  - The cron poller queries the PREVIOUS day, because you finalise returns by 11am "
        "ET, and follows the cursor to the end.\n"
        "  - A Notification of Change (tranStatus 3, or any C-code) never changes the order; "
        "we note the change code on it. We hold no account details - every debit is "
        "authorized fresh through Bank Connect - so there is nothing on our side to correct.\n"
        "  - A return on an unpaid order moves it to 'returned' and releases stock. A return on "
        "an order we had already counted as paid or shipped is flagged to a human instead. "
        "Return codes are stored per order (R01, R02, R03, R08, R10).\n"
        "  - The same report is shown in our admin console, with the meaning of tranStatus "
        "next to each row.\n"
        "\n"
        "We noticed this sandbox row says tranStatus 1 (settled) while its settlement row has "
        "settleDate null - we take that as sandbox data, as you said.\n"
        "\n"
        "REQUEST: please confirm the statement descriptor buyers will see on their bank "
        "statement. An unrecognised descriptor is the main cause of R10, and we want to stay "
        "well under the NACHA 0.5% unauthorized-return threshold.",
    ),
    (
        "SALE-10",
        "QueryReserves",
        "Pull up-to-date Reserves information",
        "No customer facing UX is expected",
        "PASS.\n"
        "Endpoint: POST /api/echeck/queryreserves\n"
        "Request: start 2026-09-01, end 2026-09-13 (re-run 2026-09-11 15:52 UTC)\n"
        "Response: HTTP 200, 1 row:\n"
        "  merchantId 83, merchantLegalName 'Ever NovA Life', debitAdjustment 100.00, "
        "debitReason '', currentReserveBalance 100.00, reserveDate 2026-09-10T18:48:38\n"
        "Page 2 (sending cursor {\"ReserveId\":34}): 0 rows, pageId null - end of report.\n"
        "\n"
        "Thank you for enabling it. Same start / end / pageId request shape as "
        "querysettlements. Understood that this row is sandbox data, not our production "
        "reserve terms. Our admin console shows this report and the current reserve balance.\n"
        "\n"
        "One check for production: our legal name is 'Ever Nova Life' (staging reads "
        "'Ever NovA Life').",
    ),
]

FOOTNOTES = [
    "",
    "Test environment",
    "Host: https://token-staging.finagy.com (staging). Merchant user 251, authenticating with "
    "the Basic Token Header you issued. api-version 3.0. paymentMethod 1 (ACH), secCode WEB.",
    "All ten results above are from live calls run between 2026-09-09 and 2026-09-11 (UTC). "
    "Nothing is simulated and nothing is taken from a mock - every transaction id can be looked "
    "up on your side.",
    "Account configuration read back from GET /api/hpp/products/ach: is_enabled true, "
    "is_email_verification_enabled false, is_user_validation_enabled true, "
    "is_user_validation_input_enabled false, rtp_validation_max_timeout_seconds 10.",
    "",
    "Transactions used for this sheet (status as of 2026-09-11 15:52 UTC)",
    "639245942322405169  ENL-CERT-S01B  $0.50   status 8 (returned R01)  - SALE-01 / SALE-07",
    "639245946244424782  ENL-CERT-S04B  $0.50   status 32 (voided)  - SALE-04",
    "ENL-CERT-S02        $1,500.00  declined at validation, no transaction created - SALE-02",
    "639245635306201861  ENL-MTU8KE4Z   $0.22   status 16 (settled)  - SALE-05, void refused",
    "639245589443869934  ENL-MTU5U6TG   $0.22   status 40 (refunded)  - SALE-06",
    "639246923014123850  (no reference) $0.22   tranCode 2 credit, status 4 (sent to bank)  "
    "- the SALE-06 refund",
    "",
    "One earlier problem, now explained",
    "Our first attempts all came back 'Transaction Declined.' and landed at status 1 "
    "(invalidated): 639245064869527503, 639245074916668658, 639245075966290768. The cause was "
    "geography, not configuration - we were testing from outside the US and Ribbit geo-blocks "
    "its bank-login portal (test.ribbit.ai and portal.ribbit.ai both answered 403). Testing "
    "from a US connection, the same chime sandbox account works first time. Real customers "
    "were never affected: we ship US-only and ACH is a US-only network.",
    "",
    "What we still need from Finagy / AllayPay",
    "1. The statement descriptor buyers will see on their bank statement. (SALE-09)",
    "2. The bank cutoff time, and whether same-day ACH is enabled on our account.",
    "3. Refund credits: can the credit carry our reference, or link back to the original "
    "debit? (SALE-06)",
    "4. A look at querysettlements paging with excludeReturnedItems=false - details in SALE-08.",
    "5. Confirmation that our production record carries the legal name 'Ever Nova Life'. "
    "(SALE-10)",
    "6. Production credentials, once this sheet is signed off.",
    "",
    "Note on manual account entry: we understand new merchants get Bank Connect only, and we "
    "are not asking to change that for launch. Worth revisiting later, since a buyer whose "
    "bank Ribbit cannot reach currently has no ACH option at all.",
]

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ACH Test Cases"

    blue = PatternFill("solid", fgColor="A4C2F4")
    cream = PatternFill("solid", fgColor="FCE8B2")
    head_font = Font(bold=True, size=11)

    ws.append(HEADERS)
    for col, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = head_font
        cell.fill = blue if col <= 4 else cream
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 32

    for row in ROWS:
        ws.append([row[0], row[1], row[2], row[3], row[4], "", ""])
        r = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left")
            cell.border = BORDER
        # the comment drives the row height: ~96 characters fit on one visual line
        lines = sum(len(part) // 96 + 1 for part in row[4].split("\n"))
        ws.row_dimensions[r].height = min(max(lines * 14, 60), 460)

    start_notes = ws.max_row + 2
    bold_headings = {
        "Test environment",
        "Transactions used for this sheet (status as of 2026-09-11 15:52 UTC)",
        "One earlier problem, now explained",
        "What we still need from Finagy / AllayPay",
    }
    for offset, text in enumerate(FOOTNOTES):
        cell = ws.cell(row=start_notes + offset, column=1, value=text)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if text in bold_headings:
            cell.font = Font(bold=True)

    widths = [12, 26, 46, 32, 96, 22, 22]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
